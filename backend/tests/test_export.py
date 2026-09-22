"""Export: the layout model, both renderers, placeholders, and the endpoints.

What can be checked here is checked against the files themselves -- the PDF is
parsed back, including by this app's own ATS parser, and the DOCX is reopened
and its XML inspected. What cannot be checked here is stated in PLAN.md rather
than implied by a green run: nobody has opened these files in Word or Google
Docs as part of this suite.
"""
from __future__ import annotations

import re

import fitz
import pytest
from docx import Document as OpenDocx
from docx.oxml.ns import qn

from app.agents.analyst.parser import parse_pdf
from app.agents.analyst.profile import CvProfile, Experience
from app.export import pdf as pdf_mod
from app.export.document import HEADINGS, build, language
from app.export.docx import _PPR, _RPR, render_docx, segments
from app.export.pdf import ExportUnavailable, is_bundled_font, render_pdf, to_html
from app.export.placeholders import PlaceholderError, clean_value, fill, slots
from tailor_fixtures import (  # noqa: F401 -- tailor_client is a fixture
    ARABIC,
    ORIGINAL,
    propose,
    tailor_client,
    upload,
)


def squash(text: str) -> str:
    return " ".join(text.split())


# --- the layout model ----------------------------------------------------------

def test_sections_follow_the_standard_order():
    keys = [s.key for s in build(ORIGINAL).sections]
    assert keys == ["summary", "experience", "education", "skills", "certifications",
                    "projects", "languages"]


def test_empty_sections_are_left_out():
    assert [s.key for s in build(ARABIC).sections] == ["summary", "experience", "skills"]


def test_headings_are_in_the_cvs_own_language():
    assert build(ORIGINAL).sections[1].heading == "Experience"
    assert build(ARABIC).sections[1].heading == HEADINGS["ar"]["experience"]


def test_language_is_decided_by_letters_not_words():
    """An Arabic CV full of tool names is still an Arabic CV."""
    mostly_tools = ARABIC.model_copy(deep=True)
    mostly_tools.skills = ["Python", "PostgreSQL", "Docker"]
    assert language(mostly_tools) == "ar"
    assert language(ORIGINAL) == "en"


def test_a_current_role_reads_present_in_each_language():
    assert "Jan 2021 – Present" in build(ORIGINAL).text
    assert "2021 – حتى الآن" in build(ARABIC).text


def test_nothing_in_the_profile_is_lost_in_layout():
    text = build(ORIGINAL).text
    everything = [
        ORIGINAL.contact.name, ORIGINAL.contact.email, ORIGINAL.contact.phone,
        ORIGINAL.summary, *ORIGINAL.skills,
        *(b for role in ORIGINAL.experience for b in role.bullets),
        *(role.company for role in ORIGINAL.experience),
        *(role.title for role in ORIGINAL.experience),
        ORIGINAL.education[0].institution, ORIGINAL.certifications[0].name,
        ORIGINAL.projects[0].description, "Redis", "Arabic (Native)",
    ]
    for item in everything:
        assert item in text, item


def test_layout_adds_nothing_but_headings_and_dates_it_formats():
    """Everything in the layout is from the profile, a heading, or a word the
    layout itself supplies. Nothing else can appear in an exported CV."""
    doc = build(ORIGINAL)
    source = ORIGINAL.all_text + "\n" + "\n".join(HEADINGS["en"].values()) + \
        "\nPresent Technologies"
    for word in re.findall(r"[\w@.+]+", doc.text):
        assert word in source, word


# --- DOCX ----------------------------------------------------------------------

@pytest.fixture(scope="module")
def english_docx(tmp_path_factory):
    path = tmp_path_factory.mktemp("docx") / "en.docx"
    path.write_bytes(render_docx(build(ORIGINAL, title="English sample")))
    return OpenDocx(str(path))


@pytest.fixture(scope="module")
def arabic_docx(tmp_path_factory):
    path = tmp_path_factory.mktemp("docx") / "ar.docx"
    path.write_bytes(render_docx(build(ARABIC, title="Arabic sample")))
    return OpenDocx(str(path))


def test_the_docx_reopens_carrying_every_line(english_docx):
    written = "\n".join(p.text for p in english_docx.paragraphs)
    for line in build(ORIGINAL).text.splitlines():
        assert squash(line) in squash(written), line


def test_the_docx_uses_real_headings_and_real_bullets(english_docx):
    styles = [p.style.name for p in english_docx.paragraphs]
    assert styles.count("Heading 1") == len(build(ORIGINAL).sections)
    assert styles.count("List Bullet") == sum(len(r.bullets) for r in ORIGINAL.experience)


def test_the_docx_has_no_tables_text_boxes_or_header_text(english_docx):
    assert english_docx.tables == []
    body = english_docx.element.body.xml
    assert "txbxContent" not in body and "w:drawing" not in body
    section = english_docx.sections[0]
    assert not any(p.text.strip() for p in section.header.paragraphs)
    assert not any(p.text.strip() for p in section.footer.paragraphs)


def test_the_docx_is_a4(english_docx):
    section = english_docx.sections[0]
    assert round(section.page_width.mm) == 210 and round(section.page_height.mm) == 297


def test_the_docx_records_title_author_and_language(english_docx, arabic_docx):
    assert english_docx.core_properties.author == "Yasmin Alamrani"
    assert english_docx.core_properties.language == "en-US"
    assert arabic_docx.core_properties.language == "ar-SA"


def test_every_arabic_paragraph_runs_right_to_left(arabic_docx):
    for paragraph in arabic_docx.paragraphs:
        assert paragraph._p.pPr.find(qn("w:bidi")) is not None, paragraph.text
    assert arabic_docx.sections[0]._sectPr.find(qn("w:bidi")) is not None


def test_an_english_docx_does_not(english_docx):
    assert not any(p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None
                   for p in english_docx.paragraphs)
    assert english_docx.sections[0]._sectPr.find(qn("w:bidi")) is None


def _runs(document):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            rtl = run._r.rPr is not None and run._r.rPr.find(qn("w:rtl")) is not None
            yield run.text, rtl


def test_arabic_runs_are_marked_rtl_and_latin_runs_are_not(arabic_docx):
    """"PostgreSQL" inside an Arabic bullet must stay a left-to-right run."""
    runs = list(_runs(arabic_docx))
    assert ("PostgreSQL", False) in [(t.strip(), r) for t, r in runs]
    for text, rtl in runs:
        if re.search(r"[؀-ۿ]", text):
            assert rtl, text
        if re.search(r"[A-Za-z]", text):
            assert not rtl, text


def test_segments_split_by_script_and_lose_nothing():
    text = "بناء واجهة تخدم 8000 تاجر باستخدام PostgreSQL"
    pieces = segments(text)
    assert "".join(p for p, _ in pieces) == text
    assert [a for _, a in pieces] == [True, False, True, False]


def test_complex_script_font_and_size_are_set(english_docx):
    """Without szCs and a cs font, Word shows Arabic at its default size."""
    rpr = english_docx.styles["Normal"].element.rPr
    assert rpr.find(qn("w:rFonts")).get(qn("w:cs")) == "Arial"
    assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "21"
    heading = english_docx.styles["Heading 1"].element.rPr
    assert heading.find(qn("w:bCs")) is not None


@pytest.mark.parametrize("which", ["english_docx", "arabic_docx"])
def test_properties_are_written_in_schema_order(which, request):
    """Word forgives misordered XML; LibreOffice and stricter parsers do not."""
    document = request.getfixturevalue(which)
    rpr_order = [qn(t) for t in _RPR]
    ppr_order = [qn(t) for t in _PPR]
    elements = list(document.element.body.iter()) + list(document.styles.element.iter())
    for element in elements:
        if element.tag == qn("w:rPr"):
            order = rpr_order
        elif element.tag == qn("w:pPr"):
            order = ppr_order
        else:
            continue
        positions = [order.index(child.tag) for child in element if child.tag in order]
        assert positions == sorted(positions), [c.tag for c in element]


# --- PDF -----------------------------------------------------------------------

@pytest.fixture(scope="module")
def weasyprint_ready():
    try:
        pdf_mod._weasyprint()
    except ExportUnavailable as exc:
        pytest.skip(str(exc))


@pytest.fixture(scope="module")
def english_pdf(weasyprint_ready, tmp_path_factory):
    path = tmp_path_factory.mktemp("pdf") / "en.pdf"
    path.write_bytes(render_pdf(build(ORIGINAL, title="English sample")))
    return path


@pytest.fixture(scope="module")
def arabic_pdf(weasyprint_ready, tmp_path_factory):
    path = tmp_path_factory.mktemp("pdf") / "ar.pdf"
    path.write_bytes(render_pdf(build(ARABIC, title="Arabic sample")))
    return path


def pdf_text(path) -> str:
    return "\n".join(page.get_text() for page in fitz.open(path))


def pdf_fonts(path) -> set[str]:
    """Embedded font names, subset prefix and hyphens removed.

    WeasyPrint writes PostScript names ("ABCDEF+Noto-Sans-Bold"); normalising
    lets the assertions name the family rather than a naming convention.
    """
    return {f[3].split("+")[-1].replace("-", "")
            for page in fitz.open(path) for f in page.get_fonts(full=True)}


def test_the_pdf_has_a_text_layer_carrying_every_line(english_pdf):
    extracted = squash(pdf_text(english_pdf))
    for line in build(ORIGINAL).text.splitlines():
        assert squash(line) in extracted, line


def test_the_pdf_carries_nothing_beyond_the_version(english_pdf):
    """Every extracted line is from the layout, less the bullet marker."""
    source = squash(build(ORIGINAL).text)
    for line in pdf_text(english_pdf).splitlines():
        line = squash(line.replace("•", ""))
        if line:
            assert line in source, line


def test_the_pdf_embeds_the_bundled_fonts_not_system_ones(english_pdf):
    fonts = pdf_fonts(english_pdf)
    assert fonts and all(name.startswith("NotoSans") for name in fonts), fonts


def test_the_pdf_has_no_images(english_pdf):
    assert sum(len(page.get_images()) for page in fitz.open(english_pdf)) == 0


def test_the_pdf_records_title_and_author(english_pdf):
    metadata = fitz.open(english_pdf).metadata
    assert metadata["title"] == "English sample"
    assert metadata["author"] == "Yasmin Alamrani"


def test_the_pdf_passes_this_apps_own_ats_checks(english_pdf):
    """The audit's parser, pointed at the export: text layer, one column, no
    tables, nothing in the header band, fonts embedded, headings recognised."""
    report = parse_pdf(english_pdf)
    assert report.has_extractable_text
    assert report.multi_column_pages == []
    assert report.table_pages == []
    assert report.header_footer_text == []
    assert report.non_embedded_fonts == []
    assert {"summary", "experience", "education", "skills", "certifications",
            "projects", "languages"} <= set(report.detected_sections)


def test_the_arabic_pdf_embeds_tajawal(arabic_pdf):
    assert any(name.startswith("Tajawal") for name in pdf_fonts(arabic_pdf))


def test_the_arabic_pdf_is_one_column_with_a_text_layer(arabic_pdf):
    report = parse_pdf(arabic_pdf)
    assert report.has_extractable_text
    assert report.primary_language == "ar"
    assert report.multi_column_pages == []
    assert report.header_footer_text == []


def test_purely_arabic_lines_extract_in_reading_order(arabic_pdf):
    """The reason the Arabic font is Tajawal. With Noto Naskh Arabic, which
    builds letters from separate dot glyphs, every one of these came back as
    junk ("ياسميOن العمرانdي")."""
    extracted = pdf_text(arabic_pdf)
    assert "ياسمين العمراني" in extracted
    assert HEADINGS["ar"]["experience"] in extracted
    assert ARABIC.summary in extracted


@pytest.mark.xfail(strict=False, reason=(
    "Known limitation: a PDF stores glyphs in visual order, and extractors "
    "rebuild a line mixing Arabic with digits or Latin text out of logical "
    "order. The DOCX keeps logical order, so it is the safer file for an ATS "
    "reading an Arabic CV. Kept as a test so a fix upstream is noticed."))
def test_a_line_mixing_arabic_and_latin_extracts_in_logical_order(arabic_pdf):
    assert ARABIC.experience[0].bullets[0] in pdf_text(arabic_pdf)


@pytest.mark.xfail(strict=False, reason=(
    "Known limitation, shared by every font tried including system Arial: "
    "a lam-alef ligature can extract with its letters swapped. The DOCX is "
    "unaffected."))
def test_a_lam_alef_ligature_extracts_as_written(arabic_pdf):
    assert "حتى الآن" in pdf_text(arabic_pdf)


def test_cv_text_cannot_become_markup():
    hostile = ORIGINAL.model_copy(deep=True)
    hostile.experience[0].bullets[0] = '<img src="http://169.254.169.254/latest">'
    hostile.summary = "<script>alert(1)</script>"
    rendered = to_html(build(hostile))
    assert "<img" not in rendered and "<script" not in rendered
    assert "&lt;img src=&quot;http://169.254.169.254/latest&quot;&gt;" in rendered


def test_the_renderer_fetches_the_bundled_fonts_and_nothing_else():
    fonts = pdf_mod.FONTS
    assert is_bundled_font((fonts / "NotoSans-Regular.ttf").as_uri())
    assert not is_bundled_font("http://169.254.169.254/latest/meta-data/")
    assert not is_bundled_font("file:///etc/passwd")
    assert not is_bundled_font(fonts.as_uri() + "/../../../../../../etc/passwd")
    assert not is_bundled_font((fonts / "OFL-NotoSans.txt").as_uri())


def test_a_refused_url_does_not_break_the_render(weasyprint_ready):
    """If markup ever did reach the renderer, a refused fetch is skipped, not fatal."""
    weasyprint = pdf_mod._weasyprint()
    fetcher = pdf_mod.fonts_only_fetcher()
    out = weasyprint.HTML(string='<p>x</p><img src="http://169.254.169.254/">',
                          url_fetcher=fetcher).write_pdf()
    assert out.startswith(b"%PDF")


# --- placeholders --------------------------------------------------------------

def with_placeholders() -> CvProfile:
    profile = ORIGINAL.model_copy(deep=True)
    profile.summary = "Backend engineer with [add number] years in payments."
    profile.experience[0].bullets[2] = "Cut run time by [add %] across [add number] jobs"
    return profile


ORIGINALS = {"summary": ORIGINAL.summary, "exp0.b2": ORIGINAL.experience[0].bullets[2]}


def test_slots_are_numbered_in_reading_order():
    found = slots(with_placeholders())
    assert [(s.index, s.target, s.placeholder) for s in found] == [
        (0, "summary", "[add number]"), (1, "exp0.b2", "[add %]"), (2, "exp0.b2", "[add number]"),
    ]


def test_a_fill_replaces_the_slot_and_records_the_figure():
    profile, supplied = fill(with_placeholders(), values={0: "6"}, drop=set(), originals=ORIGINALS)
    assert profile.summary == "Backend engineer with 6 years in payments."
    assert supplied == [{"target": "summary", "value": "6"}]


def test_two_slots_in_one_item_fill_independently():
    profile, _ = fill(with_placeholders(), values={2: "40"}, drop=set(), originals=ORIGINALS)
    assert profile.experience[0].bullets[2] == "Cut run time by [add %] across 40 jobs"


def test_a_drop_restores_the_original_wording():
    profile, supplied = fill(with_placeholders(), values={}, drop={1}, originals=ORIGINALS)
    assert profile.experience[0].bullets[2] == ORIGINAL.experience[0].bullets[2]
    assert supplied == []


def test_a_drop_wins_over_a_fill_in_the_same_item():
    profile, supplied = fill(with_placeholders(), values={2: "40"}, drop={1}, originals=ORIGINALS)
    assert profile.experience[0].bullets[2] == ORIGINAL.experience[0].bullets[2]
    assert supplied == []


def test_every_slot_settled_leaves_none():
    profile, _ = fill(with_placeholders(), values={0: "6", 2: "40"}, drop={1}, originals=ORIGINALS)
    assert slots(profile) == []


@pytest.mark.parametrize("bad", ["", "   ", "about forty", "[40]", "4" * 40])
def test_a_figure_has_to_be_a_figure(bad):
    with pytest.raises(PlaceholderError):
        clean_value(bad)


@pytest.mark.parametrize("good", ["40%", "SAR 2.5M", "1,200", "٤٠"])
def test_real_figures_are_accepted(good):
    assert clean_value(good) == good


def test_an_unknown_slot_is_refused():
    with pytest.raises(PlaceholderError, match="no placeholder 9"):
        fill(with_placeholders(), values={9: "1"}, drop=set(), originals=ORIGINALS)


def test_a_drop_needs_the_original_on_record():
    with pytest.raises(PlaceholderError, match="not on record"):
        fill(with_placeholders(), values={}, drop={0}, originals={})


# --- the endpoints -------------------------------------------------------------

def _tailored_version(client) -> dict:
    """A saved version with every honest edit accepted -- one of which leaves
    an [add %] placeholder, moved by a reorder to the top of its role."""
    resume_id = upload(client)
    proposal = propose(client, resume_id).json()
    saved = client.post(f"/api/tailor/{proposal['proposal_id']}/versions",
                        json={"accepted_ids": [e["id"] for e in proposal["edits"]]})
    assert saved.status_code == 200, saved.text
    version = saved.json()
    assert version["placeholders"] == 1
    return version


def test_export_is_refused_while_a_placeholder_remains(tailor_client):
    version = _tailored_version(tailor_client)
    for fmt in ("pdf", "docx"):
        response = tailor_client.get(f"/api/versions/{version['id']}/export.{fmt}")
        assert response.status_code == 409, fmt
        assert "placeholder" in response.json()["detail"]


def test_placeholders_are_listed_with_their_context(tailor_client):
    version = _tailored_version(tailor_client)
    found = tailor_client.get(f"/api/versions/{version['id']}/placeholders").json()
    assert len(found) == 1
    assert found[0]["placeholder"] == "[add %]"
    assert "reconciliation" in found[0]["text"]


def test_a_filled_version_exports_as_docx(tailor_client):
    version = _tailored_version(tailor_client)
    filled = tailor_client.post(f"/api/versions/{version['id']}/placeholders",
                                json={"values": {"0": "35%"}})
    assert filled.status_code == 200 and filled.json()["placeholders"] == 0

    response = tailor_client.get(f"/api/versions/{version['id']}/export.docx")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.content[:2] == b"PK"
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"


def test_a_filled_version_exports_as_pdf(tailor_client, weasyprint_ready):
    version = _tailored_version(tailor_client)
    tailor_client.post(f"/api/versions/{version['id']}/placeholders", json={"values": {"0": "35%"}})
    response = tailor_client.get(f"/api/versions/{version['id']}/export.pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")
    assert "35%" in "\n".join(p.get_text() for p in fitz.open(stream=response.content, filetype="pdf"))


def test_dropping_restores_the_original_sentence_and_unblocks_export(tailor_client):
    version = _tailored_version(tailor_client)
    dropped = tailor_client.post(f"/api/versions/{version['id']}/placeholders",
                                 json={"drop": [0]}).json()
    assert dropped["placeholders"] == 0
    bullets = CvProfile.model_validate(dropped["profile"]).experience[0].bullets
    assert ORIGINAL.experience[0].bullets[2] in bullets
    assert tailor_client.get(f"/api/versions/{version['id']}/export.docx").status_code == 200


def test_the_figures_a_user_typed_are_recorded(tailor_client):
    from app.models.db import SessionLocal
    from app.models.entities import CvVersion

    version = _tailored_version(tailor_client)
    tailor_client.post(f"/api/versions/{version['id']}/placeholders", json={"values": {"0": "35%"}})
    db = SessionLocal()
    row = db.get(CvVersion, version["id"])
    assert row.user_supplied == [{"target": "exp0.b0", "value": "35%"}]
    db.close()


def test_a_figure_that_is_not_a_figure_is_refused(tailor_client):
    version = _tailored_version(tailor_client)
    response = tailor_client.post(f"/api/versions/{version['id']}/placeholders",
                                  json={"values": {"0": "a lot"}})
    assert response.status_code == 422


def test_the_original_exports_without_any_filling(tailor_client):
    version = _tailored_version(tailor_client)
    versions = tailor_client.get(f"/api/resumes/{version['resume_id']}/versions").json()
    original = next(v for v in versions if v["is_original"])
    assert tailor_client.get(f"/api/versions/{original['id']}/export.docx").status_code == 200


def test_a_missing_pdf_renderer_is_a_503_with_a_way_forward(tailor_client, monkeypatch):
    from app.api import export as export_mod

    def unavailable(doc):
        raise ExportUnavailable("PDF export needs Pango. Word export still works.")

    monkeypatch.setattr(export_mod, "render_pdf", unavailable)
    version = _tailored_version(tailor_client)
    tailor_client.post(f"/api/versions/{version['id']}/placeholders", json={"drop": [0]})
    response = tailor_client.get(f"/api/versions/{version['id']}/export.pdf")
    assert response.status_code == 503
    assert "Word export still works" in response.json()["detail"]


def test_an_unknown_version_cannot_be_exported(tailor_client):
    assert tailor_client.get("/api/versions/nope/export.pdf").status_code == 404


def test_only_pdf_and_docx_are_formats(tailor_client):
    version = _tailored_version(tailor_client)
    assert tailor_client.get(f"/api/versions/{version['id']}/export.exe").status_code == 422


def test_an_arabic_filename_survives_in_its_encoded_form():
    from app.api.export import _filename

    full, fallback = _filename("ياسمين العمراني", "مهندسة برمجيات", "pdf")
    assert full == "ياسمين العمراني - مهندسة برمجيات.pdf"
    assert fallback.isascii() and fallback.endswith(".pdf")


def test_a_filename_cannot_carry_path_or_header_characters():
    from app.api.export import _filename

    full, fallback = _filename('../../etc/"passwd"\r\nX-Evil: 1', "", "docx")
    for bad in ("/", "\\", '"', "\r", "\n"):
        assert bad not in full and bad not in fallback
